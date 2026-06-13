from io import BytesIO

from app.domain.documents.entities import ExtractedDocument, ExtractedPage


class DocxParserAdapter:
    supported_types = {"docx"}

    def parse(self, document_id: str, filename: str, file_type: str, content: bytes) -> ExtractedDocument:
        from docx import Document

        doc = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        headings = [
            paragraph.text.strip()
            for paragraph in doc.paragraphs
            if paragraph.style and paragraph.style.name.lower().startswith("heading") and paragraph.text.strip()
        ]
        table_text = []
        tables = []
        for table_index, table in enumerate(doc.tables, start=1):
            rows = []
            for row_index, row in enumerate(table.rows, start=1):
                values = [cell.text.strip() for cell in row.cells]
                rows.append(f"Row {row_index}: " + ", ".join(values))
            tables.append({"table_number": table_index, "rows": [[cell.text.strip() for cell in row.cells] for row in table.rows]})
            table_text.append(f"Table {table_index}\n" + "\n".join(rows))
        text = "\n\n".join(paragraphs + table_text)
        return ExtractedDocument(
            document_id=document_id,
            filename=filename,
            file_type=file_type,
            pages=[ExtractedPage(page_number=1, text=text, headings=headings, tables=tables)],
            metadata={"parser": "python-docx", "paragraph_count": len(paragraphs), "table_count": len(doc.tables)},
        )

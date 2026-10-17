from io import BytesIO

from docx import Document
from pypdf import PdfWriter

from app.adapters.parsers.docx_parser_adapter import DocxParserAdapter
from app.adapters.ocr.sarvam_document_intelligence_adapter import (
    decode_output_file,
    pages_from_document_intelligence_payload,
    page_metrics,
)
from app.adapters.parsers.pdf_parser_adapter import PDFParserAdapter
from app.adapters.parsers.text_parser_adapter import TextParserAdapter
from app.application.documents.service import merge_ocr_result
from app.application.documents.service import _file_type
from app.domain.documents.entities import ExtractedDocument, ExtractedPage, OCRResult


def test_file_type_normalization():
    assert _file_type("notes.MARKDOWN") == "md"
    assert _file_type("contract.PDF") == "pdf"
    assert _file_type("no-extension") == ""


def test_text_parser_extracts_markdown_headings():
    parsed = TextParserAdapter().parse("doc-1", "notes.md", "md", b"# Overview\n\nBody")

    assert parsed.pages[0].text.startswith("# Overview")
    assert parsed.pages[0].headings == ["Overview"]


def test_docx_parser_extracts_paragraphs_and_tables():
    buffer = BytesIO()
    doc = Document()
    doc.add_heading("Payment Terms", level=1)
    doc.add_paragraph("Invoices are due within 30 days.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Delay"
    table.rows[0].cells[1].text = "Penalty"
    doc.save(buffer)

    parsed = DocxParserAdapter().parse("doc-1", "agreement.docx", "docx", buffer.getvalue())

    assert "Invoices are due within 30 days." in parsed.full_text
    assert "Payment Terms" in parsed.pages[0].headings
    assert "Penalty" in parsed.full_text
    assert parsed.pages[0].tables[0]["rows"][0] == ["Delay", "Penalty"]


def test_pdf_parser_reads_page_count():
    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)

    parsed = PDFParserAdapter().parse("doc-1", "blank.pdf", "pdf", buffer.getvalue())

    assert parsed.metadata["page_count"] == 1
    assert parsed.metadata["ocr_required_pages"] == [1]
    assert parsed.metadata["ocr_status"] == "required_not_configured"
    assert parsed.pages[0].page_number == 1


def test_document_intelligence_payload_extracts_pages():
    pages = pages_from_document_intelligence_payload(
        {
            "outputs": {
                "result.json": {
                    "pages": [
                        {"page_number": 1, "markdown": "Scanned invoice text", "tables": [{"rows": [["A", "B"]]}]},
                    ]
                }
            }
        }
    )

    assert pages[0].page_number == 1
    assert pages[0].text == "Scanned invoice text"
    assert pages[0].tables[0]["rows"][0] == ["A", "B"]


def test_decode_output_file_reads_zip_json():
    import json
    import zipfile

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("output.json", json.dumps({"pages": [{"page_number": 1, "text": "OCR text"}]}))

    decoded = decode_output_file("output.zip", buffer.getvalue())

    assert decoded["output.json"]["pages"][0]["text"] == "OCR text"


def test_page_metrics_reads_nested_job_details():
    metrics = page_metrics({"job_details": [{"total_pages": 2, "pages_processed": 2, "pages_succeeded": 1}]})

    assert metrics["total_pages"] == 2
    assert metrics["pages_succeeded"] == 1


def test_merge_ocr_result_fills_blank_pdf_pages():
    extracted = ExtractedDocument(
        document_id="doc-1",
        filename="scan.pdf",
        file_type="pdf",
        pages=[ExtractedPage(page_number=1, text="", headings=[])],
        metadata={"ocr_required_pages": [1], "ocr_status": "required_not_configured"},
    )
    result = OCRResult(
        pages=[ExtractedPage(page_number=1, text="Scanned OCR text", headings=[])],
        metadata={"ocr_provider": "fake", "ocr_status": "completed"},
        raw_payload={"pages": [{"text": "Scanned OCR text"}]},
    )

    merged = merge_ocr_result(extracted, result)

    assert merged.pages[0].text == "Scanned OCR text"
    assert merged.metadata["ocr_status"] == "completed"
    assert merged.metadata["ocr_completed_pages"] == [1]
    assert merged.metadata["_ocr_raw_payload"]["pages"][0]["text"] == "Scanned OCR text"


def _pdf_document(pages, ocr_required):
    from app.domain.documents.entities import ExtractedDocument, ExtractedPage

    return ExtractedDocument(
        document_id="doc-1",
        filename="scan.pdf",
        file_type="pdf",
        pages=[ExtractedPage(page_number=i, text=t) for i, t in enumerate(pages, start=1)],
        metadata={"ocr_required_pages": ocr_required, "ocr_status": "required_not_configured"},
    )


def _service_with_failing_ocr():
    from app.application.documents.service import DocumentService

    class FailingOCR:
        def extract_pdf(self, *args, **kwargs):
            raise RuntimeError("provider down")

    service = DocumentService.__new__(DocumentService)
    service.ocr = FailingOCR()
    return service


def test_ocr_failure_keeps_the_pages_that_already_had_text():
    from app.infrastructure.db.models import DocumentModel

    service = _service_with_failing_ocr()
    document = DocumentModel(id="doc-1", filename="scan.pdf", file_type="pdf")
    # Page 2 has no text layer and OCR is down, but page 1 is perfectly readable.
    extracted = _pdf_document(["Payment terms are 30 days.", ""], [2])

    result = service._apply_pdf_ocr_if_needed(document, b"%PDF-", extracted)

    assert result.pages[0].text == "Payment terms are 30 days."
    assert result.metadata["ocr_degraded"] is True
    assert result.metadata["ocr_status"] == "error"


def test_document_with_no_readable_text_anywhere_reports_that_specifically():
    import pytest

    from app.infrastructure.db.models import DocumentModel
    from app.shared.user_errors import DOCUMENT_NO_READABLE_TEXT

    service = _service_with_failing_ocr()
    document = DocumentModel(id="doc-1", filename="scan.pdf", file_type="pdf")
    extracted = _pdf_document(["", ""], [1, 2])

    with pytest.raises(RuntimeError) as error:
        service._apply_pdf_ocr_if_needed(document, b"%PDF-", extracted)

    assert str(error.value) == DOCUMENT_NO_READABLE_TEXT
